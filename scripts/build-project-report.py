#!/usr/bin/env python3
"""Build an editable DOCX and fixed-pagination PDF from the same authored report.

Run from the repository root after installing docs/project-report/requirements.txt.
The supplied college reference establishes visual format, not Parkline facts.
"""
from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape
from copy import deepcopy
import json
import re
import textwrap
import uuid
import zipfile

import pymupdf
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle, Image, Preformatted, Spacer
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs/project-report"
SOURCE = BASE / "source"
ASSETS = BASE / "assets"
FONTS = ASSETS / "fonts"
META = json.loads((SOURCE / "metadata.json").read_text())
PDF_PATH = BASE / "Parkline_BCA_Project_Report.pdf"
DOCX_PATH = BASE / "Parkline_BCA_Project_Report.docx"
PAGE_W, PAGE_H = A4
LEFT, RIGHT, TOP, BOTTOM = 90, 72, 72, 72
WIDTH = PAGE_W - LEFT - RIGHT
HEIGHT = PAGE_H - TOP - BOTTOM
NAVY, BLUE, LINE = "#142D4E", "#224A8A", "#D8E1EB"

for name, filename in [
    ("STIX", "STIXGeneral.ttf"), ("STIXB", "STIXGeneralBol.ttf"),
    ("STIXI", "STIXGeneralItalic.ttf"), ("STIXBI", "STIXGeneralBolIta.ttf"),
    ("Mono", "DejaVuSansMono.ttf"),
]:
    pdfmetrics.registerFont(TTFont(name, str(FONTS / filename)))
pdfmetrics.registerFontFamily("STIX", normal="STIX", bold="STIXB", italic="STIXI", boldItalic="STIXBI")

STYLES = {
    "body": ParagraphStyle("Body", fontName="STIX", fontSize=12, leading=18, alignment=TA_JUSTIFY, spaceAfter=7),
    "list": ParagraphStyle("List", fontName="STIX", fontSize=12, leading=18, leftIndent=14, firstLineIndent=-14, alignment=TA_LEFT, spaceAfter=7),
    "h2": ParagraphStyle("H2", fontName="STIXB", fontSize=13.5, leading=17.5, textColor=colors.HexColor(NAVY), spaceAfter=8),
    "h3": ParagraphStyle("H3", fontName="STIXB", fontSize=12, leading=16, textColor=colors.HexColor(NAVY), spaceBefore=5, spaceAfter=6),
    "chapter": ParagraphStyle("Chapter", fontName="STIXB", fontSize=14, leading=18, alignment=TA_CENTER, textColor=colors.HexColor(BLUE)),
    "chapter_title": ParagraphStyle("ChapterTitle", fontName="STIXB", fontSize=16, leading=21, alignment=TA_CENTER, textColor=colors.HexColor(NAVY), spaceAfter=30),
    "front_title": ParagraphStyle("FrontTitle", fontName="STIXB", fontSize=15, leading=20, alignment=TA_CENTER, textColor=colors.HexColor(NAVY), spaceAfter=26),
    "caption": ParagraphStyle("Caption", fontName="STIXB", fontSize=10.5, leading=14, spaceBefore=3, spaceAfter=8),
    "cell": ParagraphStyle("Cell", fontName="STIX", fontSize=9.3, leading=12.5, alignment=TA_LEFT),
    "cell_head": ParagraphStyle("CellHead", fontName="STIXB", fontSize=9.5, leading=12.5, textColor=colors.HexColor(NAVY)),
    "code": ParagraphStyle("Code", fontName="Mono", fontSize=7.8, leading=10.4, leftIndent=0, spaceAfter=10),
    "reference": ParagraphStyle("Reference", fontName="STIX", fontSize=10.8, leading=15.5, alignment=TA_LEFT, spaceAfter=8),
    "small": ParagraphStyle("Small", fontName="STIXI", fontSize=9.5, leading=13, spaceAfter=8),
    "center": ParagraphStyle("Center", fontName="STIX", fontSize=12, leading=18, alignment=TA_CENTER, spaceAfter=8),
    "right": ParagraphStyle("Right", fontName="STIX", fontSize=12, leading=18, alignment=TA_RIGHT, spaceAfter=8),
}

@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list = field(default_factory=list)
    path: str = ""
    anchor: str = ""
    entries: list = field(default_factory=list)
    height: float = 0

@dataclass
class Page:
    id: str
    chapter: str
    blocks: list[Block]
    number: int = 0
    front: bool = False


def inline(text: str) -> str:
    """Conservative Markdown inline formatting to ReportLab mini-XML."""
    parts = re.split(r"(`[^`]+`|\*\*.+?\*\*)", text)
    out = []
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            out.append('<font name="Mono" size="9.2">' + escape(part[1:-1]) + "</font>")
        elif part.startswith("**") and part.endswith("**"):
            out.append("<b>" + escape(part[2:-2]) + "</b>")
        else:
            out.append(escape(part))
    return "".join(out)


def parse_blocks(text: str, page_id: str) -> list[Block]:
    lines = text.strip().splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("```"):
            code = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            blocks.append(Block("code", "\n".join(code))); i += 1; continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            blocks.append(Block("heading", heading[2], len(heading[1]))); i += 1; continue
        image = re.match(r"^!\[(.+)\]\((.+)\)$", line)
        if image:
            blocks.append(Block("image", image[1], path=image[2])); i += 1; continue
        if line.startswith("**Table ") and line.endswith("**"):
            blocks.append(Block("caption", line[2:-2])); i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-+:?$", cell) for cell in row): rows.append(row)
                i += 1
            if any(len(row) != len(rows[0]) for row in rows): raise ValueError(f"Unequal table rows in {page_id}")
            blocks.append(Block("table", rows=rows)); continue
        if re.match(r"^\d+\.\s", line) or line.startswith("- "):
            blocks.append(Block("list", line)); i += 1; continue
        paragraph = [line]; i += 1
        while i < len(lines) and lines[i].strip():
            other = lines[i].strip()
            if re.match(r"^(#{1,3}\s|```|!\[|\||\*\*Table |\d+\.\s|- )", other): break
            paragraph.append(other); i += 1
        blocks.append(Block("body", " ".join(paragraph)))
    for n, block in enumerate(blocks):
        if block.kind in ["heading", "caption", "image"]: block.anchor = f"{page_id}-{n}"
    return blocks


def load_body() -> list[Page]:
    pages = []
    for source in sorted(SOURCE.glob("[0-9][0-9]-*.md")):
        data = source.read_text()
        sections = re.split(r"<!-- page:\s*([^|]+?)\s*\|\s*(.*?)\s*-->", data)
        for i in range(1, len(sections), 3):
            page_id, chapter, content = sections[i:i+3]
            blocks = parse_blocks(content, page_id.strip())
            if chapter.strip() == "REFERENCES":
                for block in blocks:
                    if block.kind == "body": block.kind = "reference"
            pages.append(Page(page_id.strip(), chapter.strip(), blocks, len(pages)+1))
    assert len({p.id for p in pages}) == len(pages), "Duplicate page IDs"
    return pages


def chapter_parts(text):
    if " — " in text: return text.split(" — ", 1)
    return "", text


def column_widths(rows):
    headers, n = rows[0], len(rows[0])
    lower = [s.lower() for s in headers]
    if n == 2: weights = [0.20, 0.80] if lower[0] == "acronym" else [0.33, 0.67]
    elif n == 3:
        if lower[-1] == "result": weights = [.12, .75, .13]
        elif lower[0] == "id" and "evidence" in lower[2]: weights = [.10, .35, .55]
        elif lower[0] == "id" or lower[0].startswith("proposed id"): weights = [.13, .50, .37]
        elif "type" in lower[1]: weights = [.30, .22, .48]
        else: weights = [.25, .36, .39]
    elif n == 4:
        if lower[0] == "id" and lower[-1] == "result": weights = [.11, .35, .40, .14]
        elif lower[0] == "id" and lower[1] == "classification": weights = [.10, .19, .34, .37]
        elif lower[0] == "id": weights = [.10, .30, .30, .30]
        elif lower[0].startswith("location"): weights = [.46, .18, .18, .18]
        elif "manifest" in lower[2]: weights = [.18, .22, .18, .42]
        elif lower[0] == "criterion": weights = [.20, .26, .27, .27]
        else: weights = [.18, .25, .29, .28]
    elif n == 5: weights = [.48, .13, .13, .13, .13]
    elif n == 6: weights = [.25, .09, .11, .11, .11, .33]
    else: weights = [1/n]*n
    return [WIDTH*w/sum(weights) for w in weights]


def wrapped_code(text):
    # Soft display wraps preserve source content and avoid right-edge clipping.
    out = []
    max_width = WIDTH - 16
    for line in text.expandtabs(2).splitlines():
        if not line:
            out.append(""); continue
        current, rest = "", line
        while pdfmetrics.stringWidth(rest, "Mono", STYLES["code"].fontSize) > max_width:
            limit = max(1, int(max_width / pdfmetrics.stringWidth("M", "Mono", STYLES["code"].fontSize)))
            cut = rest.rfind(" ", 0, limit+1)
            if cut < limit//2: cut = limit
            out.append(rest[:cut])
            rest = "    " + rest[cut:].lstrip()
        out.append(rest)
    return "\n".join(out)


def image_size(block):
    p = BASE / block.path
    with PILImage.open(p) as im: iw, ih = im.size
    max_w, max_h = WIDTH, 330
    if "/screenshots/" in block.path:
        if "overview.png" in block.path: max_h = 320
        elif "mobile-upper" in block.path: max_w, max_h = 330, 360
        elif "form" in block.path or "validation" in block.path: max_w, max_h = 330, 330
    if block.height: max_h = min(max_h, block.height)
    ratio = min(max_w/iw, max_h/ih)
    return iw*ratio, ih*ratio


def index_style(entry):
    level = entry.get("level", 2)
    return ParagraphStyle("Index", fontName="STIXB" if level == 1 else "STIX", fontSize=11.2 if level == 1 else 10.6,
                          leading=14.5 if level == 1 else 13.8, leftIndent=0 if level == 1 else (12 if level == 2 else 25),
                          textColor=colors.HexColor(NAVY) if level == 1 else colors.black)


def index_row(entry):
    label = inline(entry["label"])
    if entry.get("anchor"): label = f'<link href="#{entry["anchor"]}" color="{NAVY}">{label}</link>'
    table = Table([[Paragraph(label,index_style(entry)), Paragraph(str(entry["page"]), ParagraphStyle("IndexPage",fontName="STIX",fontSize=10.6,leading=14,alignment=TA_RIGHT))]],colWidths=[WIDTH-35,35])
    table.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),0),("TOPPADDING",(0,0),(-1,-1),4 if entry.get("level")==1 else 1.2),("BOTTOMPADDING",(0,0),(-1,-1),4 if entry.get("level")==1 else 1.2)]))
    return table


def flows_for_block(block):
    k = block.kind
    if k == "heading":
        if block.level == 1:
            prefix, title = chapter_parts(block.text)
            if prefix:
                return [(Paragraph(inline(prefix),STYLES["chapter"]),0,4),(Paragraph(inline(title),STYLES["chapter_title"]),0,30)]
            return [(Paragraph(inline(title),STYLES["chapter_title"]),0,26)]
        style = STYLES["h2" if block.level == 2 else "h3"]
        return [(Paragraph(inline(block.text),style), style.spaceBefore, style.spaceAfter)]
    if k in ["body","list","small","center","right","caption","front_title","reference"]:
        style = STYLES[k]
        return [(Paragraph(inline(block.text),style),style.spaceBefore,style.spaceAfter)]
    if k == "gap": return [(Spacer(1,block.height),0,0)]
    if k == "code":
        f = Preformatted(wrapped_code(block.text),STYLES["code"])
        return [(f,3,10)]
    if k == "signature":
        rows = [[Paragraph(inline(cell),ParagraphStyle("Signature",fontName="STIXB" if r==0 else "STIX",fontSize=11.5,leading=17)) for cell in row] for r,row in enumerate(block.rows)]
        t=Table(rows,colWidths=[WIDTH*.5,WIDTH*.5])
        t.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),0),("RIGHTPADDING",(0,0),(-1,-1),14),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
        return [(t,0,9)]
    if k == "table":
        rows = [[Paragraph(inline(cell),STYLES["cell_head" if r==0 else "cell"]) for cell in row] for r,row in enumerate(block.rows)]
        t = Table(rows,colWidths=column_widths(block.rows),repeatRows=1,hAlign="LEFT")
        t.setStyle(TableStyle([
            ("VALIGN",(0,0),(-1,-1),"TOP"), ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EAF0F7")),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7F9FB")]),
            ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#CAD5E0")),
            ("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),
            ("TOPPADDING",(0,0),(-1,-1),3.5 if block.rows[0][0]=="Acronym" else 4),("BOTTOMPADDING",(0,0),(-1,-1),3.5 if block.rows[0][0]=="Acronym" else 4),
        ]))
        return [(t,0,9)]
    if k == "image":
        w,h = image_size(block)
        image = Image(str(BASE / block.path),width=w,height=h)
        caption = Paragraph(inline(block.text),STYLES["caption"])
        return [(image,4,7),(caption,0,9)]
    if k == "index": return [(index_row(entry),0,0) for entry in block.entries]
    raise ValueError(f"Unknown block kind: {k}")


def measure_page(page):
    height = 0
    for block in page.blocks:
        for flow,before,after in flows_for_block(block):
            _,h = flow.wrap(WIDTH,10000)
            height += before+h+after
    return height


def roman(n):
    values = [(1000,"m"),(900,"cm"),(500,"d"),(400,"cd"),(100,"c"),(90,"xc"),(50,"l"),(40,"xl"),(10,"x"),(9,"ix"),(5,"v"),(4,"iv"),(1,"i")]
    result = ""
    for v,s in values:
        while n>=v: result+=s; n-=v
    return result


def front_fixed():
    m=META
    title=m["title"]
    certificate=[
        Block("front_title","BONAFIDE CERTIFICATE"),
        Block("body",f'This is to certify that the project report titled **“{title}”** is a bonafide work carried out by **{m["student_name"]} ({m["registration_number"]})** under my supervision for the award of the Degree of Bachelor of Computer Applications. To my knowledge, the work reported herein is the original work of the candidate.'),
        Block("gap",height=55),
        Block("signature",rows=[["PROJECT GUIDE","HEAD OF DEPARTMENT"],[m["guide_name"]+"; "+m["guide_designation"],m["head_name"]+"; Department of Computer Applications"],["Signature: __________________","Signature: __________________"]]),
        Block("gap",height=55),
        Block("signature",rows=[["INTERNAL EXAMINER","EXTERNAL EXAMINER"],["Signature: __________________","Signature: __________________"]]),
        Block("gap",height=30),
        Block("small","Submission template: complete candidate and signatory details and obtain actual institutional approval. No signature or approval is asserted by this generated document."),
    ]
    declaration=[
        Block("front_title","DECLARATION"),
        Block("body",f'I, **{m["student_name"]} ({m["registration_number"]})**, hereby declare that the project report entitled **“{title}”**, submitted to SRM Institute of Science and Technology in partial fulfillment of the requirements for the award of the Degree of Bachelor of Computer Applications, is a record of original software engineering and academic project work carried out by me under the supervision and guidance of **{m["guide_name"]}**, {m["guide_designation"]}, Department of Computer Applications, Directorate of Distance Education.'),
        Block("body","I further declare that this report or any part thereof has not been submitted earlier to another university or institution for the award of a degree, diploma or fellowship. Source materials, software dependencies and the supplied formatting reference are acknowledged in the report."),
        Block("gap",height=65),
        Block("right",m["student_name"]),Block("right",m["registration_number"]),
        Block("right","Signature of the Candidate: __________________"),
        Block("body",f'Place: {m["place"]}'),Block("body",f'Date: {m["submission_month"].title()}'),
        Block("small","Candidate review is required before signing. Personal details are intentionally left as editable placeholders."),
    ]
    acknowledgement=[
        Block("front_title","ACKNOWLEDGEMENT"),
        Block("body","I express my sincere gratitude to SRM Institute of Science and Technology and the Directorate of Distance Education for providing the academic opportunity to undertake this application-development project."),
        Block("body","I thank the Department of Computer Applications, the programme coordinator and the faculty members for the learning environment and project-review framework that support the development of analytical, programming and documentation skills."),
        Block("body",f'I extend my gratitude to my project guide, **{m["guide_name"]}**, for guidance, technical review and encouragement during the analysis, design, implementation and documentation of this project.'),
        Block("body","I am grateful to my family, friends and peers for their encouragement throughout the preparation of the project and its documentation. I also acknowledge the maintainers of the open-source frameworks, libraries and verification tools used by the application."),
        Block("body","The supplied college project report was used as a reference for academic structure and visual formatting. Its unrelated implementation claims and personal identities were not transferred into this report. All technical descriptions of Parkline are grounded in the reviewed code and the stated evidence boundary."),
        Block("gap",height=35),Block("right",m["student_name"]),
    ]
    abstract=[
        Block("front_title","ABSTRACT"),
        Block("body","Managing parking across multiple floors requires a consistent view of individual bays, active vehicles and planned capacity. Disconnected records can confuse reserved inventory with availability, permit duplicate assignments or lose occupied resources during capacity changes. This project presents Parkline, a browser-based multi-level parking operations and capacity planning application that makes these relationships visible and testable."),
        Block("body","The application is implemented with Next.js 16.4.3, React 19.0.0, TypeScript and a Tailwind/CSS styling pipeline. Typed fixtures provide four demonstration garage locations. A React dashboard supports location and floor selection, bay search and filtering, vehicle arrival, active-session departure, protected floor resizing and configurable new locations. Pure domain helpers enforce available-bay allocation, normalized Tamil Nadu registration input, location-local duplicate-session rejection and preservation of occupied or reserved bays during capacity reduction."),
        Block("body","Planning and operational changes are serialized into browser local storage. This is device-local demo persistence, not a shared database. The current application has no authentication, role-based authorization, payment processing, reservation lifecycle, gate control or live sensor connection. Static interface labels and sample trends are distinguished from derived occupancy data throughout the report."),
        Block("body","The reviewed revision passed nine existing unit tests, linting, type checking and a production build. A complete run of three existing browser scenarios produced two passes and one reload-persistence failure. Sixteen supplemental documentation assertions produced fifteen passes and one narrow-screen overflow failure. Genuine screenshots, source-grounded diagrams, requirements, a data dictionary and an open-findings register document both the working MVP and its limitations."),
        Block("body","The project demonstrates a practical local state model and an evidence-based engineering workflow. It provides a foundation for further hardening and later server-backed operations, while avoiding unsupported claims of production readiness or physical deployment."),
        Block("small","Keywords: Multi-Level Parking; Capacity Planning; React; Next.js; TypeScript; Browser Local Storage; State Transitions; Vehicle Sessions; Vitest; Playwright."),
    ]
    return [Page("cover","",[],front=True),Page("certificate","",certificate,front=True),Page("declaration","",declaration,front=True),Page("acknowledgement","",acknowledgement,front=True),Page("abstract","",abstract,front=True)]


ABBREVIATIONS = [
    ("A / O / R / T","Available / Occupied / Reserved / Total bay counts"),
    ("API","Application Programming Interface — no parking API in this MVP"),
    ("ARIA","Accessible Rich Internet Applications"),
    ("B1","Basement level identifier in the sample floor plan"),
    ("BCA","Bachelor of Computer Applications"),
    ("CSS","Cascading Style Sheets"),
    ("DFD","Data Flow Diagram"),
    ("DOM","Document Object Model"),
    ("E2E","End-to-End browser testing"),
    ("ER","Entity Relationship — conceptual in this report"),
    ("EV","Electric Vehicle; bay-kind indicator only"),
    ("FR / NFR","Functional Requirement / Non-Functional Requirement"),
    ("HTML","HyperText Markup Language"),
    ("HTTP / HTTPS","HyperText Transfer Protocol / HTTP Secure"),
    ("ID","Identifier; several bay and floor IDs are location-scoped"),
    ("ISO","International Organization for Standardization; date-string context"),
    ("JSON","JavaScript Object Notation"),
    ("L1 / L2","Ground-floor / upper-deck identifiers in the sample plan"),
    ("MVP","Minimum Viable Product"),
    ("npm","JavaScript package manager used by this repository"),
    ("OWASP","Open Worldwide Application Security Project"),
    ("RBAC","Role-Based Access Control — future capability, not implemented"),
    ("RTO","Regional Transport Office; syntax digits only in plate validation"),
    ("SRS","Software Requirements Specification"),
    ("TN","Tamil Nadu prefix in the project's registration policy"),
    ("TS / TSX","TypeScript / TypeScript with JSX markup"),
    ("UAT","User Acceptance Testing — no formal stakeholder UAT conducted"),
    ("UI / UX","User Interface / User Experience"),
    ("UML","Unified Modeling Language"),
    ("UTC","Coordinated Universal Time"),
    ("WCAG","Web Content Accessibility Guidelines"),
]


def split_index(entries, title):
    chunks, current, used = [], [], 0
    limit = HEIGHT-50
    for entry in entries:
        _,h=index_row(entry).wrap(WIDTH,10000)
        if current and used+h>limit:
            chunks.append(current); current=[]; used=0
        current.append(entry); used+=h
    if current: chunks.append(current)
    return chunks


def make_front(body):
    headings,figures,tables=[],[],[]
    for page in body:
        for block in page.blocks:
            entry={"label":block.text,"page":page.number,"anchor":block.anchor,"level":block.level or 2}
            if block.kind=="heading": headings.append(entry)
            elif block.kind=="image": figures.append({**entry,"level":2})
            elif block.kind=="caption": tables.append({**entry,"level":2})
    fixed=front_fixed()
    prefixes=[{"label":label,"page":roman(i+1),"anchor":p.id,"level":2} for i,(label,p) in enumerate(zip(["Title Page","Bonafide Certificate","Declaration","Acknowledgement","Abstract"],fixed))]
    more=[{"label":label,"page":"xiii","anchor":anchor,"level":2} for label,anchor in [("Table of Contents","contents-1"),("List of Figures","figures-1"),("List of Tables","tables-1"),("Abbreviations / Acronyms","abbreviations")]]
    toc_chunks=split_index(prefixes+more+headings,"TABLE OF CONTENTS")
    figure_chunks=split_index(figures,"LIST OF FIGURES")
    table_chunks=split_index(tables,"LIST OF TABLES")
    offsets={"contents-1":6,"figures-1":6+len(toc_chunks),"tables-1":6+len(toc_chunks)+len(figure_chunks),"abbreviations":6+len(toc_chunks)+len(figure_chunks)+len(table_chunks)}
    for entry in more: entry["page"]=roman(offsets[entry["anchor"]])
    toc_chunks=split_index(prefixes+more+headings,"TABLE OF CONTENTS")
    front=fixed[:]
    for stem,title,chunks in [("contents","TABLE OF CONTENTS",toc_chunks),("figures","LIST OF FIGURES",figure_chunks),("tables","LIST OF TABLES",table_chunks)]:
        for i,chunk in enumerate(chunks,1):
            front.append(Page(f"{stem}-{i}","",[Block("front_title",title+(" (CONTINUED)" if i>1 else "")),Block("index",entries=chunk)],front=True))
    front.append(Page("abbreviations","",[Block("front_title","ABBREVIATIONS / ACRONYMS"),Block("table",rows=[["Acronym","Meaning / contextual qualification"]]+[list(r) for r in ABBREVIATIONS])],front=True))
    for i,page in enumerate(front,1): page.number=i
    return front, headings, figures, tables


def draw_cover(c):
    c.setStrokeColor(colors.HexColor(NAVY)); c.setLineWidth(1.8); c.rect(27,27,PAGE_W-54,PAGE_H-54)
    c.setStrokeColor(colors.HexColor("#7F8B9E")); c.setLineWidth(.7); c.rect(32,32,PAGE_W-64,PAGE_H-64)
    center=PAGE_W/2
    def line(text,y,size=12,bold=True):
        c.setFillColor(colors.black); c.setFont("STIXB" if bold else "STIX",size); c.drawCentredString(center,y,text)
    for text,y in zip(["PARKLINE: MULTI-LEVEL PARKING","OPERATIONS AND CAPACITY","PLANNING SYSTEM"],[756,735,714]): line(text,y,15.5)
    line("A PROJECT REPORT SUBMITTED TO",676)
    line(META["institution"],652,13)
    line("IN PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE",624)
    line("AWARD OF THE DEGREE OF",608)
    line(META["degree"],573,15)
    line("BY",543)
    line(META["student_name"],519,14)
    line("REG. NO. "+META["registration_number"],499,12)
    line("UNDER THE GUIDANCE OF",467)
    line(META["guide_name"]+" "+META["guide_qualification"],443,11.5)
    line(META["guide_designation"]+", SRMIST-DDE",423,11.5)
    c.drawImage(str(ASSETS/"srm-logo.jpeg"),center-86,339,width=172,height=60,mask="auto")
    for text,y in [(META["department"],316),(META["directorate"],299),(META["institution"],282),("(Deemed to be university u/s 3 of UGC Act, 1956)",265),("Kattankulathur – 603 203",248),("Chennai, Tamil Nadu",231),(META["submission_month"],214)]: line(text,y,12)


def draw_running(c,page):
    if not page.front:
        c.setStrokeColor(colors.HexColor(LINE)); c.setLineWidth(.6); c.line(LEFT,PAGE_H-52,PAGE_W-RIGHT,PAGE_H-52)
        c.setFillColor(colors.HexColor(NAVY)); c.setFont("STIXI",8.6)
        c.drawString(LEFT,PAGE_H-46,"PARKLINE — MULTI-LEVEL PARKING SYSTEM")
        size=8.6
        while pdfmetrics.stringWidth(page.chapter,"STIXI",size)>WIDTH-195: size-=.2
        c.setFont("STIXI",size); c.drawRightString(PAGE_W-RIGHT,PAGE_H-46,page.chapter)
        c.line(LEFT,54,PAGE_W-RIGHT,54)
    if page.id!="cover":
        c.setFillColor(colors.HexColor(NAVY)); c.setFont("STIX",10.5)
        c.drawCentredString(PAGE_W/2,38,roman(page.number) if page.front else f"Page {page.number}")


def build_pdf(front,body):
    c=canvas.Canvas(str(PDF_PATH),pagesize=A4,pageCompression=1)
    c.setTitle(META["title"]); c.setAuthor("Parkline project documentation")
    c.setSubject("BCA academic project report grounded in the reviewed Parkline source")
    c.setKeywords("Parkline, parking, BCA, Next.js, TypeScript, local-storage MVP")
    c.addPageLabel(0,style="ROMAN_LOWER",start=1); c.addPageLabel(len(front),style="ARABIC",start=1)
    layout=[]
    for page in front+body:
        c.bookmarkPage(page.id)
        if page.id=="cover":
            draw_cover(c); c.showPage(); layout.append({"id":page.id,"usedHeight":0}); continue
        y=PAGE_H-TOP
        for block in page.blocks:
            if block.anchor:
                c.bookmarkPage(block.anchor)
                if block.kind=="heading" and block.level==1: c.addOutlineEntry(block.text,block.anchor,level=0,closed=False)
            for flow,before,after in flows_for_block(block):
                w,h=flow.wrap(WIDTH,10000); y-=before
                x=LEFT+(WIDTH-w)/2 if isinstance(flow,Image) else LEFT
                flow.drawOn(c,x,y-h); y-=h+after
        if y<BOTTOM-0.5: raise ValueError(f"Overflow on {page.id}: {BOTTOM-y:.1f} pt; no PDF completion claim allowed")
        draw_running(c,page); c.showPage()
        layout.append({"id":page.id,"number":roman(page.number) if page.front else page.number,"usedHeight":round(PAGE_H-TOP-y,2),"freeHeight":round(y-BOTTOM,2)})
    c.save()
    return layout


# Word rendering uses the same content blocks and explicit page starts. The PDF is
# authoritative for pagination; embedded fonts reduce reflow in compatible viewers.
def set_font(run,size=12,bold=False,italic=False,name="STIXGeneral",color=None):
    run.font.name=name; run.font.size=Pt(size); run.bold=bold; run.italic=italic
    fonts=run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ["ascii","hAnsi","eastAsia","cs"]: fonts.set(qn("w:"+attr),name)
    if color: run.font.color.rgb=RGBColor.from_string(color.lstrip("#"))


def add_inline(p,text,size=12,bold=False,italic=False,color=None):
    for part in re.split(r"(`[^`]+`|\*\*.+?\*\*)",text):
        code=part.startswith("`") and part.endswith("`")
        strong=part.startswith("**") and part.endswith("**")
        run=p.add_run(part[1:-1] if code else part[2:-2] if strong else part)
        set_font(run,9.2 if code else size,bold or strong,italic,"DejaVu Sans Mono" if code else "STIXGeneral",color)


def p_format(p,line=18,after=8,before=0,align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    f=p.paragraph_format; f.line_spacing=Pt(line); f.space_after=Pt(after); f.space_before=Pt(before); f.widow_control=False; p.alignment=align


def add_bookmark(p,name,num):
    start=OxmlElement("w:bookmarkStart"); start.set(qn("w:id"),str(num)); start.set(qn("w:name"),name)
    end=OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"),str(num))
    p._p.insert(0,start); p._p.append(end)


def configure_section(section,front=False,cover=False,start=None):
    section.page_width=Pt(PAGE_W); section.page_height=Pt(PAGE_H)
    section.left_margin=Pt(72 if cover else LEFT); section.right_margin=Pt(RIGHT)
    section.top_margin=Pt(TOP); section.bottom_margin=Pt(BOTTOM)
    section.header_distance=Pt(30); section.footer_distance=Pt(30)
    section.header.is_linked_to_previous=False; section.footer.is_linked_to_previous=False
    if not cover:
        for border in list(section._sectPr.findall(qn("w:pgBorders"))): section._sectPr.remove(border)
    pgnum=section._sectPr.find(qn("w:pgNumType"))
    if pgnum is None: pgnum=OxmlElement("w:pgNumType"); section._sectPr.append(pgnum)
    pgnum.set(qn("w:fmt"),"lowerRoman" if front else "decimal")
    if start is not None: pgnum.set(qn("w:start"),str(start))
    else: pgnum.attrib.pop(qn("w:start"),None)


def add_footer(section,front=False):
    p=section.footer.paragraphs[0]; p_format(p,line=12,after=0,align=WD_ALIGN_PARAGRAPH.CENTER)
    if not front: set_font(p.add_run("Page "),10.5)
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE")
    run=OxmlElement("w:r"); text=OxmlElement("w:t"); text.text="1"; run.append(text); field.append(run); p._p.append(field)


def add_header(section,text):
    p=section.header.paragraphs[0]; p_format(p,line=11,after=3,align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run("PARKLINE — MULTI-LEVEL PARKING SYSTEM"),8,italic=True,color=NAVY)
    set_font(p.add_run("\t"+text),7.7,italic=True,color=NAVY)
    borders=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
    for k,v in [("val","single"),("sz","4"),("space","4"),("color",LINE[1:])]: bottom.set(qn("w:"+k),v)
    borders.append(bottom); p._p.get_or_add_pPr().append(borders)


def doc_gap(doc,height):
    p=doc.add_paragraph(); p_format(p,line=max(1,height),after=0,align=WD_ALIGN_PARAGRAPH.LEFT); set_font(p.add_run(""),1)


def word_table(doc,rows,widths=None,index=False):
    table=doc.add_table(rows=0,cols=len(rows[0])); table.autofit=False
    widths=widths or column_widths(rows)
    for col,w in zip(table.columns,widths): col.width=Pt(w)
    for rid,row in enumerate(rows):
        cells=table.add_row().cells
        for cid,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Pt(widths[cid]); p=cell.paragraphs[0]
            p_format(p,line=12.5,after=0,align=WD_ALIGN_PARAGRAPH.LEFT)
            add_inline(p,text,9.5 if rid==0 else 9.3,bold=(rid==0),color=NAVY if rid==0 else None)
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement("w:tcMar")
            for side,val in [("top",70 if rows[0][0]=="Acronym" else 80),("bottom",70 if rows[0][0]=="Acronym" else 80),("left",120),("right",120)]:
                e=OxmlElement("w:"+side); e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa"); margins.append(e)
            tcpr.append(margins)
            shading=OxmlElement("w:shd"); shading.set(qn("w:fill"),"EAF0F7" if rid==0 else "F7F9FB" if rid%2==0 else "FFFFFF"); tcpr.append(shading)
        trpr=table.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement("w:cantSplit"); trpr.append(cant)
        if rid==0: trpr.append(OxmlElement("w:tblHeader"))
    borders=OxmlElement("w:tblBorders")
    for side in ["top","left","bottom","right","insideH","insideV"]:
        el=OxmlElement("w:"+side)
        for k,v in [("val","single"),("sz","3"),("color","CAD5E0")]: el.set(qn("w:"+k),v)
        borders.append(el)
    table._tbl.tblPr.append(borders)
    doc_gap(doc,9)
    return table


def word_index(doc,entries):
    for entry in entries:
        table=doc.add_table(rows=1,cols=2); table.autofit=False
        table.columns[0].width=Pt(WIDTH-35); table.columns[1].width=Pt(35)
        for i,w in enumerate([WIDTH-35,35]): table.cell(0,i).width=Pt(w)
        p=table.cell(0,0).paragraphs[0]; style=index_style(entry)
        p_format(p,line=style.leading,after=2.4 if entry.get("level")!=1 else 8,align=WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.left_indent=Pt(style.leftIndent)
        if entry.get("anchor"):
            link=OxmlElement("w:hyperlink"); link.set(qn("w:anchor"),entry["anchor"])
            run=OxmlElement("w:r"); rpr=OxmlElement("w:rPr"); fonts=OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"),"STIXGeneral"); fonts.set(qn("w:hAnsi"),"STIXGeneral"); rpr.append(fonts)
            size=OxmlElement("w:sz"); size.set(qn("w:val"),str(round(style.fontSize*2))); rpr.append(size)
            if entry.get("level")==1: rpr.append(OxmlElement("w:b"))
            run.append(rpr); t=OxmlElement("w:t"); t.text=entry["label"]; run.append(t); link.append(run); p._p.append(link)
        else: add_inline(p,entry["label"],style.fontSize,bold=entry.get("level")==1)
        p=table.cell(0,1).paragraphs[0]; p_format(p,line=14,after=0,align=WD_ALIGN_PARAGRAPH.RIGHT); set_font(p.add_run(str(entry["page"])),10.6)
        for cell in table.rows[0].cells:
            tcpr=cell._tc.get_or_add_tcPr(); mar=OxmlElement("w:tcMar")
            for side in ["top","bottom","left","right"]:
                e=OxmlElement("w:"+side); e.set(qn("w:w"),"0"); e.set(qn("w:type"),"dxa"); mar.append(e)
            tcpr.append(mar)


def word_cover(doc):
    sec=doc.sections[0]; configure_section(sec,front=True,cover=True,start=1)
    borders=OxmlElement("w:pgBorders"); borders.set(qn("w:offsetFrom"),"page")
    for side in ["top","left","bottom","right"]:
        b=OxmlElement("w:"+side)
        for k,v in [("val","double"),("sz","8"),("space","24"),("color",NAVY[1:])]: b.set(qn("w:"+k),v)
        borders.append(b)
    sec._sectPr.append(borders)
    def line(text,size=12,before=0,after=0,leading=18):
        p=doc.add_paragraph(); p_format(p,line=leading,after=after,before=before,align=WD_ALIGN_PARAGRAPH.CENTER); set_font(p.add_run(text),size,bold=True)
    line("PARKLINE: MULTI-LEVEL PARKING\nOPERATIONS AND CAPACITY\nPLANNING SYSTEM",15.5,leading=21,after=18)
    line("A PROJECT REPORT SUBMITTED TO",after=6)
    line(META["institution"],13,after=13)
    line("IN PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE\nAWARD OF THE DEGREE OF",leading=16,after=20)
    line(META["degree"],15,after=12)
    line("BY",after=6); line(META["student_name"],14,after=2)
    line("REG. NO. "+META["registration_number"],after=14)
    line("UNDER THE GUIDANCE OF",after=6)
    line(META["guide_name"]+" "+META["guide_qualification"],11.5)
    line(META["guide_designation"]+", SRMIST-DDE",11.5,after=10)
    p=doc.add_paragraph(); p_format(p,line=64,after=8,align=WD_ALIGN_PARAGRAPH.CENTER); p.add_run().add_picture(str(ASSETS/"srm-logo.jpeg"),width=Pt(172))
    for text in [META["department"],META["directorate"],META["institution"],"(Deemed to be university u/s 3 of UGC Act, 1956)","Kattankulathur – 603 203","Chennai, Tamil Nadu",META["submission_month"]]: line(text,12,leading=16)


def embed_fonts(doc):
    """Embed redistributable font resources using the OOXML obfuscation scheme."""
    part=doc.part.part_related_by(RT.FONT_TABLE)
    tree=parse_xml(part.blob)
    families=[("STIXGeneral",[("embedRegular","STIXGeneral.ttf"),("embedBold","STIXGeneralBol.ttf"),("embedItalic","STIXGeneralItalic.ttf"),("embedBoldItalic","STIXGeneralBolIta.ttf")]),("DejaVu Sans Mono",[("embedRegular","DejaVuSansMono.ttf")])]
    for family,variants in families:
        font=next((e for e in tree if e.get(qn("w:name"))==family),None)
        if font is None:
            font=OxmlElement("w:font"); font.set(qn("w:name"),family); tree.append(font)
        for tag,filename in variants:
            key=uuid.uuid4(); mask=bytes.fromhex(key.hex)[::-1]
            data=bytearray((FONTS/filename).read_bytes())
            for i in range(32): data[i]^=mask[i%16]
            font_part=Part(PackURI("/word/fonts/"+filename+".odttf"),"application/vnd.openxmlformats-officedocument.obfuscatedFont",bytes(data),doc.part.package)
            rid=part.relate_to(font_part,RT.FONT)
            emb=OxmlElement("w:"+tag); emb.set(qn("r:id"),rid); emb.set(qn("w:fontKey"),"{"+str(key).upper()+"}"); font.append(emb)
    part._blob=__import__("lxml.etree",fromlist=["etree"]).tostring(tree,xml_declaration=True,encoding="UTF-8",standalone=True)
    settings=doc.settings.element
    if settings.find(qn("w:embedTrueTypeFonts")) is None: settings.append(OxmlElement("w:embedTrueTypeFonts"))


def build_docx(front,body):
    doc=Document()
    normal=doc.styles["Normal"]; normal.font.name="STIXGeneral"; normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=Pt(18); normal.paragraph_format.space_after=Pt(8)
    for level,size in [(1,16),(2,13.5),(3,12)]:
        style=doc.styles[f"Heading {level}"]; style.font.name="STIXGeneral"; style.font.size=Pt(size); style.font.bold=True; style.font.color.rgb=RGBColor.from_string(NAVY[1:]); style.paragraph_format.keep_with_next=True
    doc.core_properties.title=META["title"]; doc.core_properties.author="Parkline project documentation"; doc.core_properties.subject="BCA academic project report"; doc.core_properties.keywords="Parkline; parking; browser-local MVP"
    word_cover(doc)
    add_bookmark(doc.paragraphs[0],"cover",0)
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE); configure_section(sec,front=True,start=2); add_footer(sec,front=True)
    bookmark_id=1
    previous_chapter=None
    for index,page in enumerate(front[1:]+body):
        break_before=False
        if index:
            if not page.front and page.chapter!=previous_chapter:
                sec=doc.add_section(WD_SECTION_START.NEW_PAGE); configure_section(sec,front=False,start=1 if page.number==1 else None); add_header(sec,page.chapter); add_footer(sec)
            else: break_before=True
        previous_chapter=page.chapter
        page_marked=False
        for block in page.blocks:
            if block.kind=="gap": doc_gap(doc,block.height); continue
            if block.kind=="signature":
                table=doc.add_table(rows=0,cols=2); table.autofit=False
                for col in table.columns: col.width=Pt(WIDTH/2)
                for rid,row in enumerate(block.rows):
                    cells=table.add_row().cells
                    for cell,text in zip(cells,row):
                        p=cell.paragraphs[0]; p_format(p,line=17,after=9,align=WD_ALIGN_PARAGRAPH.LEFT); add_inline(p,text,11.5,bold=rid==0)
                doc_gap(doc,9); continue
            if block.kind=="table":
                word_table(doc,block.rows); continue
            if block.kind=="index": word_index(doc,block.entries); continue
            if block.kind=="image":
                p=doc.add_paragraph(); p_format(p,line=12,after=7,before=4,align=WD_ALIGN_PARAGRAPH.CENTER)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.keep_with_next=True
                w,h=image_size(block); r=p.add_run(); set_font(r,1); r.add_picture(str(BASE/block.path),width=Pt(w),height=Pt(h))
                p=doc.add_paragraph(); p_format(p,line=14,after=9,align=WD_ALIGN_PARAGRAPH.LEFT); add_inline(p,block.text,10.5,bold=True)
            elif block.kind=="code":
                p=doc.add_paragraph(); p_format(p,line=10.4,after=10,before=3,align=WD_ALIGN_PARAGRAPH.LEFT)
                set_font(p.add_run(wrapped_code(block.text)),7.8,name="DejaVu Sans Mono")
            elif block.kind=="heading":
                p=doc.add_paragraph(style=f"Heading {block.level}")
                if block.level==1:
                    prefix,title=chapter_parts(block.text); p_format(p,line=21,after=30,align=WD_ALIGN_PARAGRAPH.CENTER)
                    if prefix: set_font(p.add_run(prefix+"\n"),14,bold=True,color=BLUE)
                    set_font(p.add_run(title),16,bold=True,color=NAVY)
                else:
                    style=STYLES["h2" if block.level==2 else "h3"]
                    p_format(p,line=style.leading,after=style.spaceAfter,before=style.spaceBefore,align=WD_ALIGN_PARAGRAPH.LEFT)
                    add_inline(p,block.text,style.fontSize,bold=True,color=NAVY)
            else:
                style=STYLES[block.kind]; p=doc.add_paragraph()
                alignment={TA_LEFT:WD_ALIGN_PARAGRAPH.LEFT,TA_CENTER:WD_ALIGN_PARAGRAPH.CENTER,TA_RIGHT:WD_ALIGN_PARAGRAPH.RIGHT,TA_JUSTIFY:WD_ALIGN_PARAGRAPH.JUSTIFY}[style.alignment]
                p_format(p,line=style.leading,after=style.spaceAfter,before=style.spaceBefore,align=alignment)
                add_inline(p,block.text,style.fontSize,bold=block.kind in ["front_title","caption"],italic=block.kind=="small",color=NAVY if block.kind=="front_title" else None)
                if block.kind=="list": p.paragraph_format.left_indent=Pt(14); p.paragraph_format.first_line_indent=Pt(-14)
            if block.kind=="caption": p.paragraph_format.keep_with_next=True
            if not page_marked:
                if break_before: p.paragraph_format.page_break_before=True
                add_bookmark(p,page.id,bookmark_id); bookmark_id+=1; page_marked=True
            if block.anchor:
                add_bookmark(p,block.anchor,bookmark_id); bookmark_id+=1
    update=OxmlElement("w:updateFields"); update.set(qn("w:val"),"true"); doc.settings.element.append(update)
    embed_fonts(doc)
    doc.save(DOCX_PATH)


def main():
    body=load_body(); front,headings,figures,tables=make_front(body)
    # Fit image height to the space left by its accompanying authored explanation.
    for page in body:
        images=[b for b in page.blocks if b.kind=="image"]
        if images:
            excess=measure_page(page)-HEIGHT
            if excess>0:
                image=images[0]; _,original_h=image_size(image)
                if excess<original_h*.24: image.height=original_h-excess-3
    overflow=[]
    for p in front[1:]+body:
        measured=measure_page(p)
        if measured>HEIGHT+.5: overflow.append((p.id,p.number,round(measured,1),round(measured-HEIGHT,1)))
    if overflow:
        print("LAYOUT OVERFLOW (page, number, required height, excess):")
        for row in overflow: print(row)
        raise SystemExit(2)
    layout=build_pdf(front,body); build_docx(front,body)
    with pymupdf.open(PDF_PATH) as pdf:
        assert len(pdf)>=90
        assert len(pdf)==len(front)+len(body)
        text="\n".join(p.get_text() for p in pdf)
        assert "PARKLINE" in text and "FR-18" in text and "D-01" in text
        assert "INFEARA MEDICAL SYSTEM" not in text
        out_of_bounds=[]
        for i,page in enumerate(pdf):
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines",[]):
                    for span in line["spans"]:
                        x0,y0,x1,y1=span["bbox"]
                        if x0<20 or y0<20 or x1>PAGE_W-20 or y1>PAGE_H-20: out_of_bounds.append([i+1,span["text"],span["bbox"]])
        assert not out_of_bounds,out_of_bounds
        summary={"title":META["title"],"reviewedCommit":META["source_commit"],"pdfPages":len(pdf),"frontMatterPages":len(front),"mainAndAppendixPages":len(body),"figureCount":len(figures),"tableCount":len(tables),"sourceWordLikeTokens":sum(len(p.read_text().split()) for p in SOURCE.glob("[0-9][0-9]-*.md")),"pdfExtractedWords":len(text.split()),"minimumRequestedPages":90,"pageCountVerified":True,"outOfPageTextSpans":out_of_bounds,"wordPagination":"Editable document with explicit page starts and embedded fonts; PDF is pagination authority.","layout":layout}
    with zipfile.ZipFile(DOCX_PATH) as z:
        assert "word/document.xml" in z.namelist()
        assert len([p for p in z.namelist() if p.startswith("word/fonts/")])==5
    (SOURCE/"document-validation.json").write_text(json.dumps(summary,indent=2)+"\n")
    print(json.dumps({k:v for k,v in summary.items() if k!="layout"},indent=2))
    print(f"Wrote {PDF_PATH}\nWrote {DOCX_PATH}")

if __name__=="__main__": main()
